"""
Word 加载器
"""


class WordLoader:
    """Word 文件加载器"""
    
    def load(self, file_path: str) -> str:
        """加载 Word 文件"""
        try:
            from docx import Document
        except ImportError:
            raise ImportError("请安装 python-docx: pip install python-docx")
        
        doc = Document(file_path)
        text = []
        
        for para in doc.paragraphs:
            if para.text.strip():
                text.append(para.text)
        
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                text.append(row_text)
        
        return "\n".join(text)
